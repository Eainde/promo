#!/usr/bin/env python
"""
Generate Panellist_Brief_Akshay_Dipta.docx from panellist_brief.md.

Run with the repo venv, which has python-docx:
    ../.venv/bin/python generate_brief_docx.py

panellist_brief.md is the source of truth. Edit that, then re-run this.
"""

import os
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "panellist_brief.md")
OUT = os.path.join(HERE, "Panellist_Brief_Akshay_Dipta.docx")

# False = leave a labelled empty box where each screenshot goes, so the images
# can be pasted in by hand. Flip to True to embed the files again.
EMBED_IMAGES = False

DB_BLUE = RGBColor(0x00, 0x18, 0x50)
DB_ACCENT = RGBColor(0x00, 0x53, 0x9F)
GREY = RGBColor(0x59, 0x59, 0x59)

IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")


def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_runs(paragraph, text, base_bold=False, size=None, color=None):
    """Render **bold**, *italic* and `code` spans into a paragraph."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        bold, italic, mono = base_bold, False, False
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, bold = part[2:-2], True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            part, mono = part[1:-1], True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            part, italic = part[1:-1], True
        run = paragraph.add_run(part)
        run.bold = bold
        run.italic = italic
        if mono:
            run.font.name = "Consolas"
            run.font.size = Pt((size or 10.5) - 1)
        elif size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return paragraph


def add_placeholder(doc, index, rel, caption):
    """Empty labelled box to paste a screenshot into, plus its caption."""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, "F4F6FA")
    cell.width = Inches(6.1)
    table.rows[0].height = Inches(1.6)

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(4)
    add_runs(p, "SCREENSHOT %d GOES HERE" % index, base_bold=True, size=11, color=DB_ACCENT)

    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(24)
    add_runs(p2, os.path.basename(rel), size=9, color=GREY)
    for run in p2.runs:
        run.font.name = "Consolas"

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(12)
    add_runs(cap, caption, size=8.5, color=GREY)
    for run in cap.runs:
        run.italic = True


def heading(doc, text, level):
    sizes = {1: 20, 2: 16, 3: 12.5, 4: 11}
    colors = {1: DB_BLUE, 2: DB_BLUE, 3: DB_ACCENT, 4: DB_ACCENT}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level <= 2 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    add_runs(p, text, base_bold=True, size=sizes[level], color=colors[level])
    return p


def add_table(doc, rows):
    cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    for r, row in enumerate(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            value = row[c] if c < len(row) else ""
            add_runs(p, value, base_bold=(r == 0), size=9.5)
            if r == 0:
                shade(cell, "E8EDF5")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(line):
    return bool(re.fullmatch(r"\|[\s:|-]+\|", line.strip()))


def build():
    with open(SRC, encoding="utf-8") as fh:
        lines = fh.read().split("\n")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Inches(0.85))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Inches(0.9))

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(footer, "Akshay Dipta  |  Evidence Brief for the Promotion Panel  |  September 2026",
             size=8, color=GREY)

    i = 0
    img_index = 0
    missing = []
    pending_break = False
    first_block = True

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Horizontal rule starts a new page.
        if re.fullmatch(r"-{3,}", stripped):
            pending_break = True
            i += 1
            continue

        if pending_break and not first_block:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        pending_break = False
        first_block = False

        # Image with caption.
        m = IMG_RE.match(stripped)
        if m:
            caption, rel = m.group(1), m.group(2)
            path = os.path.normpath(os.path.join(HERE, rel))
            img_index += 1
            if not EMBED_IMAGES:
                add_placeholder(doc, img_index, rel, caption)
            elif os.path.exists(path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(2)
                p.add_run().add_picture(path, width=Inches(6.1))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap.paragraph_format.space_after = Pt(12)
                add_runs(cap, caption, size=8.5, color=GREY).runs[0].italic = True
                for run in cap.runs:
                    run.italic = True
            else:
                missing.append(rel)
            i += 1
            continue

        # Table.
        if stripped.startswith("|") and i + 1 < len(lines) and is_separator(lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            add_table(doc, rows)
            continue

        # Headings.
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            heading(doc, m.group(2), len(m.group(1)))
            i += 1
            continue

        # Blockquote.
        if stripped.startswith(">"):
            quote = [stripped.lstrip("> ").strip()]
            i += 1
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.right_indent = Inches(0.35)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(10)
            add_runs(p, " ".join(quote), size=10.5, color=DB_ACCENT)
            for run in p.runs:
                run.italic = True
            continue

        # Bullet.
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, stripped[2:])
            i += 1
            continue

        # Numbered item.
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, m.group(1))
            i += 1
            continue

        # Plain paragraph.
        add_runs(doc.add_paragraph(), stripped)
        i += 1

    doc.save(OUT)
    print("wrote %s (%d screenshot %s)"
          % (OUT, img_index, "slots" if not EMBED_IMAGES else "images"))
    if missing:
        print("WARNING missing images:")
        for rel in missing:
            print("  " + rel)


if __name__ == "__main__":
    build()
