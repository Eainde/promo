"""Round 2 promotion deck. 8 slides.

Adapted from generate_ppt.py. Structure differs from round 1: About Me and
My Platform are merged into one context slide, and Documents & Validity and
Messaging are merged into one foundations slide. The two slots that frees go
to the post go-live results for the Senior Manager workflow and to the
ownership extraction workflow.

House rules: no em-dashes, en-dashes or semicolons anywhere in the deck.
No body text below 12pt (2025 panel feedback said text was unreadable).

Run with python3.12, which is the interpreter that has python-pptx here.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DB_BLUE = RGBColor(0x00, 0x1E, 0x96)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
MEDIUM_GRAY = RGBColor(0xE0, 0xE0, 0xE0)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
MID_TEXT = RGBColor(0x55, 0x55, 0x55)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)
ACCENT_BLUE = RGBColor(0x42, 0x6B, 0xB4)
PALE_BLUE = RGBColor(0xE8, 0xEC, 0xF7)
GOLD = RGBColor(0xD4, 0xA0, 0x17)
LAVENDER = RGBColor(0xCC, 0xCC, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                bold=False, color=DARK_GRAY, alignment=PP_ALIGN.LEFT,
                font_name="Calibri", italic=False):
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.italic = italic
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return box


def add_shape_box(slide, left, top, width, height, fill_color, text="",
                  font_size=14, font_color=DARK_GRAY, bold=False,
                  alignment=PP_ALIGN.CENTER, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sh = slide.shapes.add_shape(shape, Inches(left), Inches(top),
                                Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill_color
    sh.line.fill.background()
    sh.shadow.inherit = False
    if text:
        tf = sh.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.alignment = alignment
        p.space_before = Pt(4)
    return sh


def add_bullet_list(slide, left, top, width, height, items, font_size=15,
                    color=DARK_GRAY, bold=False, space=8):
    box = slide.shapes.add_textbox(Inches(left), Inches(top),
                                   Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.space_before = Pt(space)
    return box


def add_metric(slide, left, top, width, metric, label, height=1.2,
               metric_size=26, label_size=13, fill=DB_BLUE):
    box = add_shape_box(slide, left, top, width, height, fill)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = metric
    p.font.size = Pt(metric_size)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(label_size)
    p2.font.color.rgb = LAVENDER
    p2.font.name = "Calibri"
    p2.alignment = PP_ALIGN.CENTER
    return box


def slide_title(slide, text, sub=None):
    add_textbox(slide, 0.7, 0.28, 12.0, 0.7, text, font_size=30,
                bold=True, color=DB_BLUE)
    if sub:
        add_textbox(slide, 0.72, 0.98, 12.0, 0.4, sub, font_size=13,
                    bold=True, color=ACCENT_BLUE)


def before_after(slide, top, before_items, after_items, height=1.9,
                 font_size=15):
    add_shape_box(slide, 0.7, top, 5.85, 0.6, RED, "BEFORE",
                  font_size=16, font_color=WHITE, bold=True)
    add_bullet_list(slide, 0.9, top + 0.75, 5.5, height, before_items,
                    font_size=font_size)
    add_shape_box(slide, 6.9, top, 5.75, 0.6, GREEN, "AFTER",
                  font_size=16, font_color=WHITE, bold=True)
    add_bullet_list(slide, 7.1, top + 0.75, 5.4, height, after_items,
                    font_size=font_size)


# 1. TITLE
s = blank()
add_background(s, DB_BLUE)
add_textbox(s, 1, 1.6, 11.3, 1.4, "Akshay Dipta", font_size=48, bold=True,
            color=WHITE, alignment=PP_ALIGN.CENTER)
add_textbox(s, 1, 3.05, 11.3, 0.8, "VP Promotion  |  Senior Engineer",
            font_size=30, color=LAVENDER, alignment=PP_ALIGN.CENTER)
add_textbox(s, 1, 4.25, 11.3, 0.6,
            "London Cadbury Team  |  Client Life Cycle Management  |  London",
            font_size=19, color=WHITE, alignment=PP_ALIGN.CENTER)
add_shape_box(s, 5.4, 5.15, 2.5, 0.04, LAVENDER, shape=MSO_SHAPE.RECTANGLE)
add_textbox(s, 1, 5.5, 11.3, 0.5,
            "Deutsche Bank, Technology, Data and Innovation",
            font_size=17, color=LAVENDER, alignment=PP_ALIGN.CENTER)

# 2. ABOUT ME + MY PLATFORM
s = blank()
slide_title(s, "About Me, and What My Platform Does")
add_shape_box(s, 0.7, 1.05, 5.85, 0.6, DB_BLUE,
              "13 Years in Financial Technology", font_size=16,
              font_color=WHITE, bold=True)
add_bullet_list(s, 0.9, 1.8, 5.5, 1.6, [
    "JP Morgan  ·  HSBC  ·  American Express  ·  Deutsche Bank",
    "Joined Deutsche Bank in March 2023",
    "Java  ·  Spring Boot  ·  Kafka  ·  AI  ·  Microservices",
], font_size=15)
add_shape_box(s, 6.9, 1.05, 5.75, 0.6, DB_BLUE, "My Roles",
              font_size=16, font_color=WHITE, bold=True)
add_bullet_list(s, 7.1, 1.8, 5.4, 1.6, [
    "AI lead for the business unit",
    "Technical authority for messaging and AI infrastructure",
    "Component Guardian. Every design and code approval goes through me",
], font_size=15)
add_shape_box(s, 0.7, 3.55, 11.95, 0.85, PALE_BLUE,
              "Before the bank can take on a client, it has to prove it knows who that client really is. "
              "Who owns them, who runs them, and whether the paperwork stands up.",
              font_size=16, font_color=DARK_GRAY, bold=True, alignment=PP_ALIGN.CENTER)
pipe = ["Document\narrives", "Linked to the right\ncompliance question",
        "Validity\nchecked", "Ready for\nregulatory review"]
xs = [0.74, 3.79, 6.84, 9.89]
for x, label in zip(xs, pipe):
    add_shape_box(s, x, 4.6, 2.7, 1.4, LIGHT_GRAY, label,
                  font_size=15, font_color=DARK_GRAY, bold=True)
for x in [3.44, 6.49, 9.54]:
    add_textbox(s, x, 4.98, 0.35, 0.6, "→", font_size=22,
                bold=True, color=ACCENT_BLUE, alignment=PP_ALIGN.CENTER)
add_textbox(s, 0.7, 6.3, 11.95, 0.6,
            "My work spans the full pipeline, and now the AI that reads the documents.",
            font_size=15, italic=True, color=ACCENT_BLUE)

# 3. AI PLATFORM
s = blank()
slide_title(s, "The AI Platform I Built",
            "SINGLE-HANDEDLY RESEARCHED, DESIGNED AND DELIVERED")
add_shape_box(s, 0.7, 1.5, 3.85, 0.55, DB_BLUE, "Nexus AI Framework",
              font_size=15, font_color=WHITE, bold=True)
add_bullet_list(s, 0.85, 2.18, 3.6, 2.45, [
    "The unit had no AI capability. Every team started from a blank page",
    "Team leaned to Python. I steered us to Java and convinced management",
    "Any team now builds a working agent in about an hour",
], font_size=13)
add_shape_box(s, 4.75, 1.5, 3.85, 0.55, DB_BLUE, "Nexus AI Studio",
              font_size=15, font_color=WHITE, bold=True)
add_bullet_list(s, 4.9, 2.18, 3.6, 2.45, [
    "A live picture of every workflow, that lights up as it runs",
    "Full prompt management. Version, compare side by side, test and publish the AI instructions, with no software release",
    "One dependency and one setting. Nothing for a team to build",
], font_size=13)
add_shape_box(s, 8.8, 1.5, 3.85, 0.55, DB_BLUE, "PromptLint Quality Gate",
              font_size=15, font_color=WHITE, bold=True)
add_bullet_list(s, 8.95, 2.18, 3.6, 2.45, [
    "Catches badly written AI instructions before they reach the model",
    "Like a spell checker for what we tell the AI to do",
    "Runs on every build, and costs nothing to run",
], font_size=13)
add_metric(s, 0.7, 4.85, 3.85, "A day  ➔  An hour", "Time to build an agent",
           metric_size=23)
add_metric(s, 4.75, 4.85, 3.85, "Presented at RCP level",
           "140 attended my session. Recorded, and still used", metric_size=23)
add_metric(s, 8.8, 4.85, 3.85, "Two live workflows",
           "Both running on this framework", metric_size=23)
add_textbox(s, 0.7, 6.3, 11.95, 0.8,
            "Multiple teams across RCP now build on this framework. But a platform nobody uses is worth nothing, so I went and used it too. The next two slides are what came out.",
            font_size=15, italic=True, color=ACCENT_BLUE)

# 4. SENIOR MANAGER WORKFLOW, AFTER GO-LIVE
s = blank()
slide_title(s, "Finding the Senior Managers Who Run a Client",
            "FIRST AI WORKFLOW IN PRODUCTION. I DESIGNED IT AND WROTE EVERY PROMPT.  |  ON MY PLATFORM")
before_after(s, 1.40, [
    "Analysts read document after document by hand to find who runs a client",
    "Each name checked across several systems, then tested against policy",
    "77 minutes of manual work per case, and rework between preparers and reviewers",
], [
    "Agents identify the people, cross-check the records and file the evidence",
    "Analysts still review the evidence and make the final decision",
    "Built in five months with Operations, Policy, Data, Controls and Transformation",
], height=1.75)
add_metric(s, 0.7, 4.05, 3.85, "77 min  ➔  Under 10",
           "More than seven times faster", height=1.25, metric_size=23)
add_metric(s, 4.75, 4.05, 3.85, "Over 99.5%", "Right first time", height=1.25, metric_size=26)
add_metric(s, 8.8, 4.05, 3.85, "57% cheaper",
           "€19.87 to €8.51 per case", height=1.25, metric_size=26)
add_shape_box(s, 0.7, 5.45, 11.95, 1.7, PALE_BLUE)
add_textbox(s, 0.95, 5.52, 11.5, 0.29,
            "Demonstrated to the national media at the bank's Bank on Tech showcase in Bengaluru, India, June 2026. Covered in the national press.",
            font_size=12.5, bold=True, color=DB_BLUE)
add_textbox(s, 0.95, 5.83, 11.5, 0.29,
            "Published bank-wide on the internal network, August 2026.",
            font_size=12.5, bold=True, color=DB_BLUE)
add_textbox(s, 0.95, 6.14, 11.5, 0.56,
            "“This is the first production step towards an AI-augmented, end to end process that is faster, more accurate "
            "and more scalable, while preserving clear human accountability for every risk decision.”",
            font_size=14, italic=True, color=DARK_GRAY)
add_textbox(s, 0.95, 6.72, 11.5, 0.33,
            "Ross Mackenzie, Co-Head of Operations and Controls, Corporate Bank and Investment Bank",
            font_size=13, bold=True, color=DB_BLUE)

# 5. OWNERSHIP AND CONTROL
s = blank()
slide_title(s, "Finding the Ownership Hierarchy Behind a Client",
            "THE HARDER ONE. I OWN IT END TO END.  |  SECOND AI WORKFLOW IN PRODUCTION  |  LIVE FOR 100 USERS")
before_after(s, 1.40, [
    "An analyst traces ownership up through layers of holding companies and trusts",
    "Documents in several languages, group policy plus every country's own rules",
    "Percentages worked out by hand. Slow, and hard to evidence to a regulator",
], [
    "One call returns the full structure, with a document quote behind every claim",
    "Collapsed an earlier six-agent pipeline into a single call",
    "Never guesses. Missing evidence is recorded as a gap, not filled in",
], height=1.75)
add_metric(s, 0.7, 4.05, 3.85, "Maker and checker",
           "A second agent scores every result", height=1.25, metric_size=22)
add_metric(s, 4.75, 4.05, 3.85, "265 automated tests",
           "I treat AI instructions as code", height=1.25, metric_size=22)
add_metric(s, 8.8, 4.05, 3.85, "Same case, same answer",
           "Measured it first, then fixed it", height=1.25, metric_size=22)
add_shape_box(s, 0.7, 5.45, 5.85, 1.7, LIGHT_GRAY)
add_textbox(s, 0.9, 5.56, 5.5, 0.32, "Partnership with the policy team",
            font_size=14, bold=True, color=DB_BLUE)
add_bullet_list(s, 0.9, 5.9, 5.5, 1.15, [
    "Audited their rules clause by clause. Put 15 written questions back",
    "Told them their own worked examples were wrong. They corrected them",
], font_size=13, space=5)
add_shape_box(s, 6.9, 5.45, 5.75, 1.7, LIGHT_GRAY)
add_textbox(s, 7.1, 5.56, 5.4, 0.32, "Built for the regulator, not just the user",
            font_size=14, bold=True, color=DB_BLUE)
add_bullet_list(s, 7.1, 5.9, 5.4, 1.15, [
    "Every assertion carries a verbatim quote from a document",
    "Versioned releases. A change is a new version, never an edit in place",
], font_size=13, space=5)

# 6. FOUNDATIONS
s = blank()
slide_title(s, "What Both of Those Workflows Stand On",
            "DOCUMENTS AND STATE. DESIGNED, BUILT AND STILL OWNED IN PRODUCTION")
add_shape_box(s, 0.7, 1.5, 7.5, 0.6, DB_BLUE,
              "Documents and the State Machine", font_size=17,
              font_color=WHITE, bold=True)
add_bullet_list(s, 0.9, 2.25, 7.2, 2.4, [
    "Analysts linked documents to compliance questions by hand. I automated it",
    "I automated the unlinking too, when a newer document supersedes an older one. "
    "A missing document is visible. A stale one still looks current to an auditor",
    "My state machine works out where every question and every document stands",
    "Both AI workflows only see work it released, and only trust documents it marked good",
], font_size=15)
add_shape_box(s, 8.4, 1.5, 4.25, 0.6, DB_BLUE,
              "Messaging", font_size=17,
              font_color=WHITE, bold=True)
add_bullet_list(s, 8.6, 2.25, 3.95, 2.4, [
    "The old infrastructure could not cope. Analysts waited 30 to 60 minutes",
    "I designed a new messaging library, including automatic retry when a message fails",
    "The same work now completes in seconds",
], font_size=15)
add_metric(s, 0.7, 4.85, 2.85, "€5 million", "Savings contributed to",
           height=1.25, metric_size=21)
add_metric(s, 3.75, 4.85, 2.85, "40,000+", "Document operations automated",
           height=1.25, metric_size=21)
add_metric(s, 6.8, 4.85, 2.85, "30-60 min ➔ Seconds", "Analyst wait time",
           height=1.25, metric_size=18)
add_metric(s, 9.85, 4.85, 2.8, "500,000 a day", "Messages processed",
           height=1.25, metric_size=21)
add_textbox(s, 0.7, 6.35, 11.95, 0.6,
            "I do not build and walk away. I was resolving production issues in the document platform this month.",
            font_size=15, italic=True, color=ACCENT_BLUE)

# 7. BEYOND MY TEAM
s = blank()
slide_title(s, "Beyond My Team",
            "REUSE, CULTURE AND CONTRIBUTION OUTSIDE THE DAY JOB")
cols = [
    (0.7, 4.0, "Industry Contribution", [
        "Contributed a new module to LangChain4j, the open source Java AI framework",
        "The same framework our AI platform is built on",
        "Giving back to the community we depend on",
    ]),
    (4.85, 4.0, "Driving Re-use", [
        "Messaging and retry libraries adopted across the business unit",
        "Nexus AI adopted unit-wide, now reviewed bank-wide",
        "Shared coding standards, now the default in every service",
    ]),
    (9.0, 3.65, "Building the Culture", [
        "2024 Deutsche Bank Hackathon, participant",
        "2025 Deutsche Bank Hackathon, team lead",
        "2026 environmental sustainability initiative",
        "Mentor engineers through pairing and code review",
    ]),
]
for left, width, head, items in cols:
    add_shape_box(s, left, 1.55, width, 0.6, DB_BLUE, head,
                  font_size=17, font_color=WHITE, bold=True)
    add_bullet_list(s, left + 0.2, 2.35, width - 0.4, 3.1, items, font_size=14)
add_shape_box(s, 0.7, 5.8, 11.95, 1.35, PALE_BLUE)
add_textbox(s, 0.95, 6.0, 11.5, 1.0,
            "Two years ago the business unit had no AI capability at all. Today it has two workflows in production, "
            "both running on a framework I built, and one of them is on the front page of the bank's internal news.",
            font_size=16, bold=True, color=DB_BLUE)

# 8. SPONSORS
s = blank()
add_textbox(s, 0.7, 0.4, 11.95, 0.9, "Thank You", font_size=42, bold=True,
            color=DB_BLUE, alignment=PP_ALIGN.CENTER)
add_shape_box(s, 5.4, 1.35, 2.5, 0.04, ACCENT_BLUE, shape=MSO_SHAPE.RECTANGLE)
add_textbox(s, 0.7, 1.5, 11.95, 0.45,
            "Akshay Dipta  |  Senior Engineer  |  London Cadbury Team, London",
            font_size=17, color=MID_TEXT, alignment=PP_ALIGN.CENTER)
add_textbox(s, 0.7, 2.0, 11.95, 0.4, "Sponsors and Endorsements",
            font_size=16, bold=True, color=DB_BLUE, alignment=PP_ALIGN.CENTER)

sponsors = [
    ("Marco Luebbers", "Managing Director",
     ["Head of Know Your Customer Operations Steering and Perimeter Governance",
      "Presented the cost and savings analysis for this work to management",
      "(To be confirmed)"]),
    ("Tong Su", "Director, CTO Initiative Lead",
     ["“One of the most outstanding software engineers I've had the pleasure to collaborate with. "
      "Consistent focus on building clean and reusable code that benefits not just our immediate project, "
      "but also the wider engineering community.”"]),
    ("Lalitha Lalwani", "Director, Architect, CTO Platform",
     ["Worked with me directly on the AI initiative",
      "Observed end to end ownership of the AI framework, from research to production"]),
    ("Chris Ashley", "Director, Conduct and Control Risk",
     ["Corporate Bank and Investment Bank, Operations and Controls",
      "Worked with me on the initial build of the documents and validity platform"]),
]
xs = [0.55, 3.70, 6.85, 10.00]
for x, (name, role, lines) in zip(xs, sponsors):
    add_shape_box(s, x, 2.55, 2.95, 4.4, LIGHT_GRAY)
    add_textbox(s, x + 0.15, 2.7, 2.65, 0.4, name, font_size=17, bold=True,
                color=DB_BLUE)
    add_textbox(s, x + 0.15, 3.1, 2.65, 0.6, role, font_size=13, bold=True,
                color=ACCENT_BLUE)
    add_bullet_list(s, x + 0.15, 3.75, 2.65, 3.0, lines, font_size=12,
                    color=DARK_GRAY, space=6)

import os
_out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "promotion_presentation_second_round.pptx")
prs.save(_out)
print("wrote %s, %d slides" % (_out, len(prs.slides._sldIdLst)))
