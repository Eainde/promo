from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Deutsche Bank colors
DB_BLUE = RGBColor(0x00, 0x1E, 0x96)
DB_DARK = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF5, 0xF5, 0xF5)
MEDIUM_GRAY = RGBColor(0xE0, 0xE0, 0xE0)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)
GREEN = RGBColor(0x2E, 0x7D, 0x32)
RED = RGBColor(0xC6, 0x28, 0x28)
ACCENT_BLUE = RGBColor(0x42, 0x6B, 0xB4)
GOLD = RGBColor(0xD4, 0xA0, 0x17)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_background(slide, color=WHITE):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, font_size=18,
                bold=False, color=DARK_GRAY, alignment=PP_ALIGN.LEFT,
                font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_shape_box(slide, left, top, width, height, fill_color, text="",
                  font_size=14, font_color=DARK_GRAY, bold=False,
                  alignment=PP_ALIGN.CENTER):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    if text:
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size = Pt(font_size)
        p.font.color.rgb = font_color
        p.font.bold = bold
        p.alignment = alignment
        tf.paragraphs[0].space_before = Pt(6)
    return shape


def add_metric_box(slide, left, top, width, metric, label):
    box = add_shape_box(slide, left, top, width, 1.2, DB_BLUE)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.add_paragraph()
    p.text = metric
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = label
    p2.font.size = Pt(13)
    p2.font.color.rgb = RGBColor(0xCC, 0xCC, 0xFF)
    p2.font.bold = False
    p2.alignment = PP_ALIGN.CENTER


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=DARK_GRAY, bold=False):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top),
                                      Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.bold = bold
        p.font.name = "Calibri"
        p.space_before = Pt(8)
        p.level = 0
    return txBox


# ─────────────────────────────────────────────
# SLIDE 1: Title Slide
# ─────────────────────────────────────────────
slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_background(slide1, DB_BLUE)

add_textbox(slide1, 1, 1.5, 11, 1.5,
            "Akshay Dipta", font_size=48, bold=True, color=WHITE,
            alignment=PP_ALIGN.CENTER)

add_textbox(slide1, 1, 3.0, 11, 0.8,
            "VP Promotion — Senior Engineer", font_size=32,
            bold=False, color=RGBColor(0xCC, 0xCC, 0xFF),
            alignment=PP_ALIGN.CENTER)

add_textbox(slide1, 1, 4.2, 11, 0.6,
            "Sieve Team  |  Client Life Cycle Management  |  London",
            font_size=20, bold=False, color=WHITE,
            alignment=PP_ALIGN.CENTER)

add_textbox(slide1, 1, 5.5, 11, 0.5,
            "Deutsche Bank — Technology, Data and Innovation",
            font_size=16, bold=False, color=RGBColor(0x99, 0x99, 0xCC),
            alignment=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# SLIDE 2: About Me
# ─────────────────────────────────────────────
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide2, WHITE)

add_textbox(slide2, 0.8, 0.4, 10, 0.8,
            "About Me", font_size=36, bold=True, color=DB_BLUE)

# Left column — Experience
add_shape_box(slide2, 0.8, 1.5, 5.5, 1.2, LIGHT_GRAY,
              "12 Years in Financial Technology", font_size=22,
              font_color=DB_BLUE, bold=True)

add_bullet_list(slide2, 1.0, 2.9, 5.0, 2.5, [
    "JP Morgan  ·  HSBC  ·  American Express  ·  Deutsche Bank",
    "Joined Deutsche Bank: March 2023",
    "Java  ·  Spring Boot  ·  Kafka  ·  AI/ML  ·  Microservices",
], font_size=17)

# Right column — Roles at DB
add_shape_box(slide2, 6.8, 1.5, 5.5, 1.2, LIGHT_GRAY,
              "My Roles at Deutsche Bank", font_size=22,
              font_color=DB_BLUE, bold=True)

add_bullet_list(slide2, 7.0, 2.9, 5.0, 2.5, [
    "Technical Lead — Sieve Team",
    "Component Guardian — Kafka & Application Health",
    "Scrum Master — Sprint Delivery & Stakeholder Comms",
    "AI Innovation Lead — Nexus AI Framework",
], font_size=17)


# ─────────────────────────────────────────────
# SLIDE 3: My Platform
# ─────────────────────────────────────────────
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide3, WHITE)

add_textbox(slide3, 0.8, 0.4, 10, 0.8,
            "My Platform — Know Your Customer (KYC)",
            font_size=34, bold=True, color=DB_BLUE)

add_textbox(slide3, 0.8, 1.3, 11, 0.8,
            "Client Life Cycle Management is the team that runs KYC for Deutsche Bank — "
            "verifying who our clients are, checking their documents, and ensuring the bank "
            "meets its regulatory obligations.",
            font_size=20, color=DARK_GRAY)

# Flow diagram: Document Upload → Auto-Association → Validity Check → Compliance Ready
flow_items = [
    ("Document\nUpload", 0.8),
    ("Auto-\nAssociation", 3.6),
    ("Validity\nCheck", 6.4),
    ("Compliance\nReady", 9.2),
]

for label, left in flow_items:
    color = DB_BLUE if label != "Compliance\nReady" else GREEN
    box = add_shape_box(slide3, left, 2.8, 2.4, 1.6, color,
                        label, font_size=18, font_color=WHITE, bold=True)

# Arrows between boxes
for i in range(3):
    left = flow_items[i][1] + 2.5
    add_textbox(slide3, left, 3.2, 0.8, 0.8, "→", font_size=36,
                bold=True, color=DB_BLUE, alignment=PP_ALIGN.CENTER)

# Description below
add_textbox(slide3, 0.8, 5.0, 11, 1.5,
            "The heart of KYC compliance — every client document flows through this platform. "
            "My work spans the full pipeline: from how documents are linked to compliance questions, "
            "to how their validity is calculated, to the messaging infrastructure that connects it all.",
            font_size=17, color=DARK_GRAY)


# ─────────────────────────────────────────────
# SLIDE 4: Bet 1 — Intelligent Document & Validity Platform
# ─────────────────────────────────────────────
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide4, WHITE)

add_textbox(slide4, 0.8, 0.3, 11, 0.8,
            "Intelligent Document & Validity Platform",
            font_size=34, bold=True, color=DB_BLUE)

add_textbox(slide4, 0.8, 1.0, 5, 0.4,
            "TEAM & PLATFORM IMPACT", font_size=14, bold=True,
            color=ACCENT_BLUE)

# BEFORE column
add_shape_box(slide4, 0.8, 1.6, 5.5, 0.7, RED,
              "BEFORE", font_size=20, font_color=WHITE, bold=True)

add_bullet_list(slide4, 1.0, 2.5, 5.0, 2.0, [
    "Analysts manually linked documents to questions",
    "Manual validity checks — slow, error-prone",
    "Client onboarding bottlenecked on document verification",
    "Fragmented legacy system, hard to maintain",
], font_size=15, color=DARK_GRAY)

# AFTER column
add_shape_box(slide4, 6.8, 1.6, 5.5, 0.7, GREEN,
              "AFTER", font_size=20, font_color=WHITE, bold=True)

add_bullet_list(slide4, 7.0, 2.5, 5.0, 2.0, [
    "Auto-Association: documents linked automatically",
    "Auto De-Association: outdated links removed",
    "Validity Framework: real-time rule checking",
    "Modular, scalable architecture I designed",
], font_size=15, color=DARK_GRAY)

# Metrics bar at bottom
add_metric_box(slide4, 0.8, 5.3, 3.5, "£25M", "Saved through automation")
add_metric_box(slide4, 4.8, 5.3, 3.5, "100s of Hours", "Analyst time freed up")
add_metric_box(slide4, 8.8, 5.3, 3.5, "Faster", "Client onboarding")


# ─────────────────────────────────────────────
# SLIDE 5: Bet 2 — Kafka Infrastructure Revolution
# ─────────────────────────────────────────────
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide5, WHITE)

add_textbox(slide5, 0.8, 0.3, 11, 0.8,
            "Kafka Infrastructure Revolution",
            font_size=34, bold=True, color=DB_BLUE)

add_textbox(slide5, 0.8, 1.0, 5, 0.4,
            "BUSINESS UNIT IMPACT — ADOPTED ACROSS CLM", font_size=14,
            bold=True, color=ACCENT_BLUE)

# BEFORE column
add_shape_box(slide5, 0.8, 1.6, 5.5, 0.7, RED,
              "BEFORE", font_size=20, font_color=WHITE, bold=True)

add_bullet_list(slide5, 1.0, 2.5, 5.0, 2.0, [
    "Users waiting 30-60 minutes for compliance results",
    "Failed messages stuck in database — no automated retry",
    "No way to trace messages through the system",
    "Daily processing lags delaying business operations",
], font_size=15, color=DARK_GRAY)

# AFTER column
add_shape_box(slide5, 6.8, 1.6, 5.5, 0.7, GREEN,
              "AFTER", font_size=20, font_color=WHITE, bold=True)

add_bullet_list(slide5, 7.0, 2.5, 5.0, 2.0, [
    "Phoenix Retry Library: intelligent, automated retry",
    "New streaming library: everything in seconds",
    "Full message observability: end-to-end tracing",
    "Component Guardian — I own this infrastructure",
], font_size=15, color=DARK_GRAY)

# Metrics bar at bottom
add_metric_box(slide5, 0.8, 5.3, 3.5, "30-60 min → Seconds",
               "Processing time reduction")
add_metric_box(slide5, 4.8, 5.3, 3.5, "500,000 msgs/day",
               "Daily message throughput")
add_metric_box(slide5, 8.8, 5.3, 3.5, "Multiple Teams",
               "Adopted across business unit")


# ─────────────────────────────────────────────
# SLIDE 6: Bet 3 — AI Agentic Platform
# ─────────────────────────────────────────────
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide6, WHITE)

add_textbox(slide6, 0.8, 0.3, 11, 0.8,
            "AI Agentic Platform — Pioneering AI Adoption",
            font_size=34, bold=True, color=DB_BLUE)

add_textbox(slide6, 0.8, 1.0, 8, 0.4,
            "INNOVATION & FUTURE — SINGLE-HANDEDLY DESIGNED & BUILT",
            font_size=14, bold=True, color=ACCENT_BLUE)

# Three-part flow: Framework → Quality → Production
add_shape_box(slide6, 0.8, 1.7, 3.5, 2.5, DB_BLUE,
              "", font_size=16, font_color=WHITE)
box1_title = add_textbox(slide6, 0.9, 1.8, 3.3, 0.5,
                         "Nexus AI Framework", font_size=20, bold=True,
                         color=WHITE, alignment=PP_ALIGN.CENTER)
add_bullet_list(slide6, 1.0, 2.4, 3.1, 1.5, [
    "Two-tier agentic architecture",
    "Declarative agent configuration",
    "Built-in guardrails & tools",
], font_size=13, color=WHITE)

add_textbox(slide6, 4.4, 2.5, 0.6, 0.6, "→", font_size=36,
            bold=True, color=DB_BLUE, alignment=PP_ALIGN.CENTER)

add_shape_box(slide6, 5.0, 1.7, 3.5, 2.5, ACCENT_BLUE)
box2_title = add_textbox(slide6, 5.1, 1.8, 3.3, 0.5,
                         "PromptLint", font_size=20, bold=True,
                         color=WHITE, alignment=PP_ALIGN.CENTER)
add_bullet_list(slide6, 5.2, 2.4, 3.1, 1.5, [
    "8-dimension quality analyzer",
    "No AI calls — runs in milliseconds",
    "CI integrated, JUnit API",
], font_size=13, color=WHITE)

add_textbox(slide6, 8.6, 2.5, 0.6, 0.6, "→", font_size=36,
            bold=True, color=DB_BLUE, alignment=PP_ALIGN.CENTER)

add_shape_box(slide6, 9.2, 1.7, 3.5, 2.5, GREEN)
box3_title = add_textbox(slide6, 9.3, 1.8, 3.3, 0.5,
                         "CSM Workflow", font_size=20, bold=True,
                         color=WHITE, alignment=PP_ALIGN.CENTER)
add_bullet_list(slide6, 9.4, 2.4, 3.1, 1.5, [
    "First production AI workflow",
    "Auto-identifies Client Senior Managers",
    "Replaces manual document reading",
], font_size=13, color=WHITE)

# Metrics bar at bottom
add_metric_box(slide6, 0.8, 5.3, 3.5, "18-24 hrs → 1 hr",
               "Agent creation time (95% reduction)")
add_metric_box(slide6, 4.8, 5.3, 3.5, "First in CLM",
               "Production AI workflow")
add_metric_box(slide6, 8.8, 5.3, 3.5, "Single-handedly",
               "Research · POC · Design · Development")


# ─────────────────────────────────────────────
# SLIDE 7: Beyond My Team
# ─────────────────────────────────────────────
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide7, WHITE)

add_textbox(slide7, 0.8, 0.4, 10, 0.8,
            "Beyond My Team", font_size=34, bold=True, color=DB_BLUE)

# Firmwide column
add_shape_box(slide7, 0.8, 1.5, 5.5, 0.8, DB_BLUE,
              "Firmwide Contributions", font_size=20,
              font_color=WHITE, bold=True)

add_bullet_list(slide7, 1.0, 2.5, 5.0, 2.5, [
    "2024: Deutsche Bank Hackathon — Participant",
    "2025: Deutsche Bank Hackathon — Team Lead",
    "2026: Environmental Sustainability Initiative (August)",
], font_size=17, color=DARK_GRAY)

# Division-wide column
add_shape_box(slide7, 6.8, 1.5, 5.5, 0.8, DB_BLUE,
              "Division-Wide Impact", font_size=20,
              font_color=WHITE, bold=True)

add_bullet_list(slide7, 7.0, 2.5, 5.0, 2.5, [
    "Kafka Phoenix Retry Library → adopted across CLM",
    "Nexus AI Framework → adopted across CLM",
    "Maven Mixin for code quality → adopted across CLM",
    "Component Guardian: Kafka infrastructure for the business unit",
], font_size=17, color=DARK_GRAY)

# Sponsors section
add_shape_box(slide7, 0.8, 5.3, 11.7, 1.2, LIGHT_GRAY,
              "Sponsors: [To be confirmed]", font_size=18,
              font_color=DARK_GRAY, bold=False,
              alignment=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# SLIDE 8: Thank You
# ─────────────────────────────────────────────
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
add_background(slide8, DB_BLUE)

add_textbox(slide8, 1, 2.0, 11, 1.5,
            "Thank You", font_size=52, bold=True, color=WHITE,
            alignment=PP_ALIGN.CENTER)

add_textbox(slide8, 1, 3.8, 11, 0.8,
            "Akshay Dipta  |  Senior Engineer  |  Sieve Team, London",
            font_size=22, bold=False, color=RGBColor(0xCC, 0xCC, 0xFF),
            alignment=PP_ALIGN.CENTER)

add_textbox(slide8, 1.5, 5.0, 10, 0.8,
            "Excited to continue driving innovation and delivering value "
            "at Deutsche Bank.",
            font_size=18, bold=False, color=WHITE,
            alignment=PP_ALIGN.CENTER)


# ─────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────
output_path = "/Users/akshaydipta/Documents/promotion/promotion_presentation.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
print(f"Total slides: {len(prs.slides)}")
